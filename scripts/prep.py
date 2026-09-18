# -*- coding: utf-8 -*-
"""
Gera data/dados.json a partir do relatório de contas a pagar (Imobiliar)
cruzado com a planilha de conferência de atrasados.

Uso:
    python scripts/prep.py

Coloque os arquivos-fonte na pasta ./fontes (não versionada):
    fontes/relatorio_ctapag_MICHEL.csv   (export do Imobiliar, ; e latin-1)
    fontes/SSDD.xlsx                      (conferência: Planilha1 + Planilha2)

Ajuste HOJE e os nomes dos arquivos abaixo se necessário.
"""
from openpyxl import load_workbook
import pandas as pd, re, json, os
from collections import Counter

# ---- configuração ----
HOJE      = pd.Timestamp('2026-09-09')          # data de referência p/ calcular atraso
DIR       = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONTES    = os.path.join(DIR, 'fontes')
BASE      = os.path.join(FONTES, 'relatorio_ctapag_MICHEL.csv')
CONF_XLSX = os.path.join(FONTES, 'SSDD.xlsx')
SAIDA     = os.path.join(DIR, 'data', 'dados.json')

# ---- conferência ----
wb = load_workbook(CONF_XLSX, read_only=True, data_only=True)
def norm(x):
    if x is None: return ''
    if hasattr(x, 'strftime'): return x.strftime('%d/%m/%Y')
    return str(x).strip()
conf = {}
for r in list(wb['Planilha1'].iter_rows(values_only=True))[1:]:
    if all(c is None for c in r): continue
    im = norm(r[0])
    if not im: continue
    deb = norm(r[2]).upper().replace('NÃO', 'NAO')
    status = 'Confirmado' if deb == 'SIM' else ('Sem débito' if deb == 'NAO'
             else ('Em progresso' if 'PROGRESSO' in deb else (deb.title() or 'Sem conferência')))
    conf[im] = {'status': status, 'adm': norm(r[1]), 'data': norm(r[3]), 'meio': norm(r[4]),
                'boleto': norm(r[5]), 'tipo': norm(r[6]), 'valor': norm(r[7]), 'quempaga': norm(r[8]),
                'motivo': norm(r[9]), 'pego': norm(r[10]), 'pago': norm(r[11]), 'dtpgto': ''}
for r in list(wb['Planilha2'].iter_rows(values_only=True))[1:]:
    if all(c is None for c in r): continue
    im = norm(r[0]); dtp = norm(r[8]) if len(r) > 8 else ''
    if im in conf and dtp: conf[im]['dtpgto'] = dtp

# ---- base consolidada ----
df = pd.read_csv(BASE, sep=';', encoding='latin-1', dtype=str).fillna('')
ORIG = {'I': 'Imóvel', 'A': 'Administrativo', 'R': 'Repasse'}
df['origem'] = df['tipoorigem'].map(ORIG).fillna(df['tipoorigem'])
df['valor_num'] = df['valor'].apply(lambda s: round((lambda t: float(t) if re.match(r'^-?\d+(\.\d+)?$', t) else 0.0)(str(s).strip().replace('.', '').replace(',', '.')), 2))
df['dt'] = pd.to_datetime(df['data'], format='%d/%m/%Y', errors='coerce')
df['comp'] = df['dt'].dt.strftime('%Y-%m')
df['venc_iso'] = df['dt'].dt.strftime('%Y-%m-%d')
df['prop'] = df['nome'].str.extract(r'Prop\. Titular\s*:\s*(\d+)')[0].fillna('')
df['ocup'] = df['nome'].str.extract(r'Status:\s*(\w+)')[0].fillna('')
df['cc'] = df['nome'].str.extract(r'^(\d+)\s*-')[0].fillna('')
df['dias'] = (HOJE - df['dt']).dt.days
def faixa(d):
    if pd.isna(d): return '—'
    d = int(d)
    if d < 0: return 'A vencer'
    if d <= 30: return '01-30 dias'
    if d <= 60: return '31-60 dias'
    if d <= 90: return '61-90 dias'
    if d <= 120: return '91-120 dias'
    if d <= 180: return '121-180 dias'
    return '+180 dias'
df['faixa'] = df['dias'].apply(faixa)
def clean(r):
    n = re.split(r'Prop\. Titular', r['nome'])[0]
    n = re.split(r'Nr\. Docum', n)[0]
    n = re.split(r'Cod\.\s*\d+', n)[0]
    return re.sub(r'\s+', ' ', n).strip(' -')
df['nome_limpo'] = df.apply(clean, axis=1)
df['parcela'] = df['parcela'].apply(lambda p: str(p).replace('"', '').strip())

EMPTY = {'status': '-', 'adm': '', 'data': '', 'meio': '', 'boleto': '', 'tipo': '',
         'valor': '', 'quempaga': '', 'motivo': '', 'pego': '', 'pago': '', 'dtpgto': ''}
def get_conf(row):
    if row['tipoorigem'] != 'I': return EMPTY
    return conf.get(str(row['codigo']), {**EMPTY, 'status': 'Sem conferência'})

recs = []
for _, r in df.iterrows():
    c = get_conf(r)
    recs.append({
        'origem': r['origem'], 'comp': r['comp'] or '-', 'venc': r['data'] or '-', 'vi': r['venc_iso'] or '',
        'dias': int(r['dias']) if pd.notna(r['dias']) else None, 'faixa': r['faixa'],
        'forn': (r['fornecedor'] or '-').strip(), 'fav': (r['favorecido'] or '-').strip(),
        'taxa': (r['taxa'] or '-').strip(), 'imovel': r['codigo'] if r['tipoorigem'] == 'I' else '',
        'nome': r['nome_limpo'][:90], 'prop': r['prop'] or '', 'ocup': r['ocup'] or '',
        'cc': r['cc'] if r['tipoorigem'] == 'A' else '', 'forma': (r['formapagto'] or '-').strip(),
        'parc': r['parcela'] or '', 'pr': r['prevreal'] or '', 'nr': r['nrlancto'],
        'compl': (r['complemento'] or '').strip()[:80], 'valor': r['valor_num'],
        'cstatus': c['status'], 'cadm': c['adm'], 'cdata': c['data'], 'cmeio': c['meio'],
        'cboleto': c['boleto'], 'ctipo': c['tipo'], 'cvalor': c['valor'], 'cquem': c['quempaga'],
        'cmotivo': c['motivo'], 'cpego': c['pego'], 'cpago': c['pago'], 'cdtpgto': c['dtpgto'],
    })

meta = {'gerado_em': HOJE.strftime('%d/%m/%Y'), 'hoje': HOJE.strftime('%d/%m/%Y'),
        'periodo_min': df['dt'].min().strftime('%d/%m/%Y'), 'periodo_max': df['dt'].max().strftime('%d/%m/%Y'),
        'total_reg': len(recs), 'total_valor': round(df['valor_num'].sum(), 2), 'conf_total': len(conf)}
os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
json.dump({'meta': meta, 'registros': recs, 'conferencia': conf}, open(SAIDA, 'w', encoding='utf-8'), ensure_ascii=False)

print('OK ->', SAIDA)
print('Registros:', len(recs), '| Total R$ {:,.2f}'.format(meta['total_valor']))
print('Origem:', dict(Counter(r['origem'] for r in recs)))
print('Conferência (I):', dict(Counter(r['cstatus'] for r in recs if r['origem'] == 'Imóvel')))
